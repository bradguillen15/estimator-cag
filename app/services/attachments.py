"""Attachment text extraction (path B: local extraction, provider-independent).

PDF (``pypdf``) and Word ``.docx`` (``python-docx``) files are turned into plain text and appended
to the transcript under a ``--- attachment: <name> ---`` separator. Because the attachments become
part of the description, they go through the same input guardrails (prompt injection, PII) and the
same cache keys as the transcript, and any model in ``LLM_MODELS`` can read them.

Known limits: images and diagrams are dropped, and scanned PDFs (pictures without a text layer)
are rejected rather than silently estimated without their content.
"""

import io
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import PurePath

import structlog
from docx import Document
from pypdf import PdfReader

from app.exceptions import AttachmentError

logger = structlog.get_logger()

MAX_ATTACHMENTS = 5
MAX_ATTACHMENT_BYTES = 5 * 1024 * 1024
# Extracted text of all the attachments together; bounds the token bill of one request.
MAX_ATTACHMENTS_CHARS = 60_000

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"  # .docx is a zip package


@dataclass(frozen=True)
class Attachment:
    filename: str
    content: bytes


def build_description(transcript: str, attachments: Sequence[Attachment]) -> str:
    """The transcript followed by the text of each attachment, under a clear separator.

    Raises ``AttachmentError`` on too many/too large files, an unsupported or unreadable file,
    a file without extractable text, or too much extracted text overall.
    """
    if len(attachments) > MAX_ATTACHMENTS:
        raise AttachmentError(f"Puedes adjuntar como máximo {MAX_ATTACHMENTS} archivos.")

    sections: list[str] = []
    total_chars = 0
    for attachment in attachments:
        name = _display_name(attachment.filename)
        text = extract_text(name, attachment.content)
        total_chars += len(text)
        if total_chars > MAX_ATTACHMENTS_CHARS:
            raise AttachmentError(
                f"Los adjuntos superan los {MAX_ATTACHMENTS_CHARS:,} caracteres de texto. "
                "Adjunta solo los documentos relevantes."
            )
        sections.append(f"--- attachment: {name} ---\n{text}")
        logger.info("attachment_extracted", filename=name, chars=len(text))

    return "\n\n".join([transcript, *sections])


def extract_text(filename: str, content: bytes) -> str:
    """Plain text of one PDF or ``.docx`` file (type checked by extension *and* magic bytes)."""
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise AttachmentError(
            f"«{filename}» supera el tamaño máximo de {MAX_ATTACHMENT_BYTES // (1024 * 1024)} MB."
        )
    extractor = _extractor_for(filename, content)
    try:
        text = extractor(content)
    except AttachmentError:
        raise
    except Exception as exc:
        # Parsers raise many types on corrupt input; the user only needs to know the file failed.
        logger.warning("attachment_unreadable", filename=filename, error_type=type(exc).__name__)
        raise AttachmentError(f"No se pudo leer «{filename}». Comprueba que no esté dañado.") from exc

    text = _normalize(text)
    if not text:
        raise AttachmentError(
            f"«{filename}» no contiene texto extraíble (¿es un PDF escaneado?). "
            "Adjunta una versión con texto."
        )
    return text


def _extractor_for(filename: str, content: bytes) -> Callable[[bytes], str]:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf" and content.startswith(_PDF_MAGIC):
        return _pdf_text
    if suffix == ".docx" and content.startswith(_ZIP_MAGIC):
        return _docx_text
    raise AttachmentError(f"«{filename}» no es un PDF ni un documento Word (.docx) válido.")


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise AttachmentError("Los PDF protegidos con contraseña no están soportados.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _display_name(filename: str) -> str:
    """Base name on a single line, so a crafted name cannot fake a separator or a new section."""
    name = PurePath(filename.replace("\\", "/")).name
    return re.sub(r"[\r\n\t]+", " ", name).strip() or "adjunto"


def _normalize(text: str) -> str:
    """Trims trailing spaces and collapses runs of blank lines left by the extractors."""
    lines = (line.rstrip() for line in text.replace("\x00", "").splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
